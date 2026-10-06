"""Load recipients from CSV or Excel. Every column becomes a template variable;
the e-mail column is found by header name. Several addresses in one cell
(separated by ; , or whitespace) go into the To header of that single mail."""

import csv
import re
from pathlib import Path

from .errors import RecipientError
from .models import Recipient

EMAIL_HEADERS = {"email", "e-mail", "mail", "邮箱", "邮件", "电子邮箱", "电子邮件", "收件人邮箱"}
_SPLIT = re.compile(r"[;,；，\s]+")
_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def load_recipients(path, sheet=None):
    return recipients_from_table(read_table(path, sheet))


def read_table(path, sheet=None):
    """Raw rows (list of lists of str) of a CSV/Excel file, header row included."""
    path = Path(path)
    if not path.is_file():
        raise RecipientError(f"找不到收件人名单: {path}")
    suffix = path.suffix.lower()
    if suffix in (".csv", ".txt"):
        rows = _read_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _read_excel(path, sheet)
    else:
        raise RecipientError(f"不支持的名单格式 {suffix}，请使用 .csv 或 .xlsx")
    return rows


def _read_csv(path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "gbk"):  # Excel on Chinese Windows saves CSV as GBK
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise RecipientError(f"无法识别 CSV 编码，请另存为 UTF-8: {path}")
    return [list(r) for r in csv.reader(text.splitlines())]


def _read_excel(path, sheet):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RecipientError("读取 Excel 需要安装 openpyxl:  pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        try:
            ws = wb[sheet] if sheet else wb.active
        except KeyError:
            raise RecipientError(f"Excel 里没有名为 {sheet!r} 的工作表，现有: {wb.sheetnames}")
        return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()  # read-only mode keeps the file handle (and a Windows lock) open


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_emails(text, what="邮箱"):
    """'a@x.com; b@x.com' -> ('a@x.com', 'b@x.com'); raises on malformed addresses."""
    emails = tuple(e for e in _SPLIT.split(str(text).strip()) if e)
    bad = [e for e in emails if not _EMAIL.match(e)]
    if bad:
        raise RecipientError(f"{what}格式不正确: {', '.join(bad)}")
    return emails


def recipients_from_table(rows, label=lambda n: f"第 {n} 行"):
    """rows[0] is the header. `label(n)` names the n-th physical row in errors."""
    numbered = [(n, r) for n, r in enumerate(rows, start=1) if any(str(c).strip() for c in r)]
    if not numbered:
        raise RecipientError("收件人名单是空的")
    headers = [str(h).strip() for h in numbered[0][1]]
    email_cols = [i for i, h in enumerate(headers) if h.lower() in EMAIL_HEADERS]
    if not email_cols:
        raise RecipientError(
            f"名单里没有邮箱列。表头需要是以下之一: {', '.join(sorted(EMAIL_HEADERS))}；当前表头: {headers}"
        )
    email_col = email_cols[0]

    recipients, problems, seen = [], [], {}
    for line_no, row in numbered[1:]:
        row = list(row) + [""] * (len(headers) - len(row))
        fields = {h: str(row[i]).strip() for i, h in enumerate(headers) if h}
        emails = tuple(e for e in _SPLIT.split(str(row[email_col]).strip()) if e)
        where = label(line_no)
        if not emails:
            problems.append(f"{where}: 邮箱为空")
            continue
        bad = [e for e in emails if not _EMAIL.match(e)]
        if bad:
            problems.append(f"{where}: 邮箱格式不正确 {', '.join(bad)}")
            continue
        key = tuple(sorted(e.lower() for e in emails))
        if key in seen:
            problems.append(f"{where}: 与{seen[key]}的收件人重复 {', '.join(emails)}")
            continue
        seen[key] = where
        recipients.append(Recipient(emails=emails, fields=fields, source=where))

    if problems:
        raise RecipientError("收件人名单有问题:\n  " + "\n  ".join(problems))
    return recipients
